from __future__ import annotations

import hashlib
import io
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from memory_policy import checked_record
from memory_store import MemoryStore

HEADERS = ["Tenant", "Entity", "Account", "Currency", "Narrative", "Vendor Hint", "Date", "Amount"]
POLICY_FIELDS = HEADERS[:6]
TEXT_FIELDS = ("Tenant", "Entity", "Account", "Currency", "Narrative")
REVIEW_SHEET = "tieout review"
OUTPUT_HEADERS = ["Record ID", "Source SHA256", "Source Sheet", "Source Cell", "Decision", "Proposed Vendor", "Reason", "Rule IDs"]


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _records_from_bytes(data: bytes, path: Path) -> list[dict]:
    sha = hashlib.sha256(data).hexdigest()
    wb = openpyxl.load_workbook(io.BytesIO(data))
    try:
        if "Transactions" not in wb.sheetnames:
            raise ValueError(f"workbook lacks 'Transactions' sheet: {path}")
        ws = wb["Transactions"]
        max_col = ws.max_column
        header_row = [ws.cell(row=1, column=col).value for col in range(1, max_col + 1)]
        columns = {}
        for header in HEADERS:
            matches = [i + 1 for i, value in enumerate(header_row) if value == header]
            if not matches:
                raise ValueError(f"missing required header '{header}' in {path}")
            if len(matches) > 1:
                raise ValueError(f"duplicate required header '{header}' in {path}")
            columns[header] = matches[0]
        narrative_letter = get_column_letter(columns["Narrative"])
        populated: dict[int, dict[int, object]] = {}
        for (r, c), cell in ws._cells.items():
            if r > 1:
                populated.setdefault(r, {})[c] = cell
        records = []
        for row in sorted(populated):
            cellmap = populated[row]
            if all(_blank(cell.value) for cell in cellmap.values()):
                continue
            values = {}
            for header in POLICY_FIELDS:
                cell = cellmap.get(columns[header])
                if cell is not None and cell.data_type == "f":
                    raise ValueError(f"formula in policy field '{header}' at {path} row {row}")
                values[header] = cell.value if cell is not None else None
            for header in TEXT_FIELDS:
                value = values[header]
                if not isinstance(value, str) or not value.strip():
                    raise ValueError(f"policy field '{header}' must be non-empty text at {path} row {row}")
                values[header] = value.strip()
            hint = values["Vendor Hint"]
            if hint is None:
                hint = ""
            if not isinstance(hint, str):
                raise ValueError(f"policy field 'Vendor Hint' must be text at {path} row {row}")
            records.append(checked_record({
                "record_id": f"{sha}:Transactions:{row}",
                "scope": {
                    "tenant": values["Tenant"],
                    "entity": values["Entity"],
                    "account": values["Account"],
                    "currency": values["Currency"],
                },
                "narrative": values["Narrative"],
                "vendor_hint": hint.strip(),
                "source": {"workbook_sha256": sha, "sheet": "Transactions", "cell": f"{narrative_letter}{row}"},
            }))
        return records
    finally:
        wb.close()


def load_records(path: Path) -> list[dict]:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"workbook must be .xlsx: {path}")
    return _records_from_bytes(path.read_bytes(), path)


def export_suggestions(source: Path, destination: Path, store: MemoryStore) -> list[dict]:
    source = Path(source)
    destination = Path(destination)
    if destination.suffix.lower() != ".xlsx":
        raise ValueError(f"export destination must be .xlsx: {destination}")
    if source.resolve() == destination.resolve():
        raise ValueError("export destination must differ from source")
    if destination.exists():
        raise FileExistsError(f"export destination already exists: {destination}")
    data = source.read_bytes()
    records = _records_from_bytes(data, source)
    suggestions = store.suggest(records)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    try:
        if REVIEW_SHEET in wb.sheetnames:
            raise ValueError(f"workbook already has '{REVIEW_SHEET}' sheet")
        ws = wb.create_sheet(REVIEW_SHEET)
        rows = [OUTPUT_HEADERS]
        for suggestion in suggestions:
            record = suggestion["record"]
            rows.append([
                record["record_id"],
                record["source"]["workbook_sha256"],
                record["source"]["sheet"],
                record["source"]["cell"],
                suggestion["decision"],
                suggestion["vendor"] or "",
                suggestion["reason"],
                " ".join(suggestion["rule_ids"]),
            ])
        for r, row in enumerate(rows, 1):
            for c, value in enumerate(row, 1):
                cell = ws.cell(row=r, column=c)
                cell.value = str(value)
                cell.data_type = "s"
        with destination.open("xb") as stream:
            wb.save(stream)
    finally:
        wb.close()
    return suggestions
