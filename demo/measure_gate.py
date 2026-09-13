"""Measure the counterparty read path on real rows, with no model in the loop.

Read-only: opens the source workbook and runs the shipped `_SCAN_CYPHER` against Aura.
Writes nothing, mutates no artifact, costs no inference credits. This is the evidence
behind the held-out gate table, the `ltd -> limited` fold counterfactual and the
"no cross-period vendor-memory axis" decision in docs/NEO4J.md.

The benchmark fixture is the first 15 staging rows with a pulled name (see
build_fixtures.py), so held-out is every row after those.

    research/.venv/bin/python demo/measure_gate.py
    research/.venv/bin/python demo/measure_gate.py --source /path/to/datasets
"""

import argparse
import os
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "harness"))

import graph  # noqa: E402

DEFAULT_SRC = Path.home() / "Downloads" / "Ylookup Hackathon Datasets"
REF = "01-bank-statements-to-journal-entries/workbook/Bank statement to journal entries - working file (anonymised).xlsx"
FIXTURE_ROWS = 15
_CCY = re.compile(
    r"\s*-\s*(usd|eur|gbp|chf|sek|dkk|nok|pln|cad|aud|jpy)\s*$", re.IGNORECASE
)

GATES = [
    ("normalised exactness", None),
    ("exactness or score >= 0.90", 0.90),
    ("exactness or score >= 0.60", 0.60),
]
CLASSES = [
    "string-exact",
    "folded-equal",
    "right entity, other spelling",
    "miss",
    "correct blank",
    "fill where golden is blank",
    "WRONG ENTITY",
]


def _entity(name) -> str:
    """Identity with currency and jurisdiction suffixes removed, then folded."""
    if not isinstance(name, str):
        return ""
    text = name.strip()
    for _ in range(3):
        candidate = graph.split_vendor(_CCY.sub("", text).strip())[0]
        if candidate == text:
            break
        text = candidate
    return graph.norm_key(text)


def _fold(value, legal_form) -> str:
    """norm_key with a swappable legal-form map, for the fold counterfactual."""
    if not isinstance(value, str):
        return ""
    ascii_only = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    tokens = [t for t in graph._FOLD.split(ascii_only.lower()) if t]
    return "".join(legal_form.get(t, t) for t in tokens)


def _classify(fill, golden) -> str:
    has_golden = bool(golden and str(golden).strip())
    if fill is None:
        return "correct blank" if not has_golden else "miss"
    if not has_golden:
        return "fill where golden is blank"
    written, want = str(fill).strip(), str(golden).strip()
    if written == want:
        return "string-exact"
    if graph.norm_key(written) == graph.norm_key(want):
        return "folded-equal"
    if _entity(written) and _entity(written) == _entity(want):
        return "right entity, other spelling"
    return "WRONG ENTITY"


def _rows(src: Path) -> list[tuple[int, str, object]]:
    wb = openpyxl.load_workbook(src, data_only=False)
    try:
        ws = wb["Staging Sheet"]
        out = []
        for r in range(2, ws.max_row + 1):
            pulled = ws.cell(r, 10).value  # Pulled Out Sender/Beneficiary
            if pulled:
                out.append((r, str(pulled).strip(), ws.cell(r, 11).value))
        return out
    finally:
        wb.close()


def _candidates(names: list[str], k: int = 3) -> dict[str, list]:
    from neo4j import RoutingControl

    queries = [
        {"raw": s, "needle": s.lower(), "norm": graph.norm_key(s)} for s in names
    ]
    result = graph._DRIVER.execute_query(
        graph._SCAN_CYPHER,
        parameters_={"queries": queries, "k": k},
        database_=graph._DB,
        routing_=RoutingControl.READ,
    )
    return {rec["pulled"]: (rec.get("cands") or []) for rec in result.records}


def _gate(cands: dict, pulled: str, min_score=None):
    """The shipped write gate: an [exact] hit (or a score floor), else blank."""
    best = (cands.get(pulled) or [None])[0]
    if best is None:
        return None
    admitted = bool(best.get("exact")) or (
        min_score is not None and (best.get("score") or 0) >= min_score
    )
    return (best.get("name") or best.get("clean")) if admitted else None


def _tally(subset, cands, min_score) -> tuple[int, dict]:
    fills, counts = 0, {}
    for _, pulled, golden in subset:
        fill = _gate(cands, pulled, min_score)
        fills += fill is not None
        key = _classify(fill, golden)
        counts[key] = counts.get(key, 0) + 1
    return fills, counts


def _print_tally(label, fills, counts, width) -> None:
    body = " | ".join(f"{k} {counts[k]}" for k in CLASSES if counts.get(k))
    print(f"  {label:<{width}} fills {fills:3d} :: {body}")


def measure_gates(rows, cands) -> None:
    held, all_rows = rows[FIXTURE_ROWS:], rows
    for label, subset in (("HELD-OUT", held), ("ALL REAL ROWS", all_rows)):
        with_golden = sum(1 for _, _, g in subset if g and str(g).strip())
        print(f"\n=== {label}: {len(subset)} rows, {with_golden} carrying a golden ===")
        for gate_label, min_score in GATES:
            fills, counts = _tally(subset, cands, min_score)
            _print_tally(gate_label, fills, counts, 30)


def measure_fold(rows, vendors) -> None:
    """The only difference between arms is _LEGAL_FORM, applied to both sides."""
    arms = {"with fold": graph._LEGAL_FORM, "no fold": {}}
    tables = {
        arm: [
            {
                _fold(s, legal_form)
                for s in [v.get("name"), v.get("clean"), *(v.get("variants") or [])]
                if s
            }
            for v in vendors
        ]
        for arm, legal_form in arms.items()
    }
    for label, subset in (("HELD-OUT", rows[FIXTURE_ROWS:]), ("ALL REAL ROWS", rows)):
        print(f"\n=== fold counterfactual, {label} ({len(subset)} rows) ===")
        for arm, table in tables.items():
            fills = right = variant = wrong = 0
            for _, pulled, golden in subset:
                want = str(golden).strip() if golden and str(golden).strip() else ""
                needle = _fold(pulled, arms[arm])
                hit = next(
                    (
                        vendors[i]
                        for i, keys in enumerate(table)
                        if needle and needle in keys
                    ),
                    None,
                )
                if hit is None:
                    continue
                fills += 1
                name = str(hit.get("name") or "").strip()
                if not want:
                    continue
                if _fold(name, graph._LEGAL_FORM) == _fold(want, graph._LEGAL_FORM):
                    right += 1
                elif _entity(name) and _entity(name) == _entity(want):
                    variant += 1
                else:
                    wrong += 1
            print(
                f"  {arm:<10} fills {fills:3d} | right identity {right:3d} | "
                f"other spelling {variant:3d} | wrong entity {wrong:3d}"
            )


def measure_recurrence(rows) -> None:
    """Whether a cross-period memory axis has anything to learn here."""
    goldens: dict[str, set] = {}
    for _, pulled, golden in rows:
        if golden and str(golden).strip():
            goldens.setdefault(graph.norm_key(pulled), set()).add(str(golden).strip())
    conflicts = {k: sorted(v) for k, v in goldens.items() if len(v) > 1}
    print(f"\n=== recurrence over {len(rows)} rows ===")
    print(f"  distinct pulled names: {len({n for _, n, _ in rows})}")
    print(f"  names resolved to more than one golden: {len(conflicts)}")
    for key, values in sorted(conflicts.items()):
        print(f"    {key} -> {values}")

    best = None
    for split in range(FIXTURE_ROWS, len(rows) - 5):
        prior = {
            graph.norm_key(n): str(g).strip()
            for _, n, g in rows[:split]
            if g and str(g).strip()
        }
        tail = [
            (graph.norm_key(n), str(g).strip() if g and str(g).strip() else "")
            for _, n, g in rows[split:]
        ]
        repeat = sum(1 for key, _ in tail if key in prior)
        agrees = sum(1 for key, g in tail if key in prior and g and g == prior[key])
        if best is None or repeat > best[1]:
            best = (split, repeat, agrees)
    split, repeat, agrees = best
    print(f"  best two-period split: {split}/{len(rows) - split}")
    print(
        f"  period-2 rows repeating a resolved name: {repeat}, prior value == golden: {agrees}"
    )


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--source",
        type=Path,
        default=Path(os.environ.get("YLOOKUP_DATASETS", DEFAULT_SRC)),
        help="dataset root holding 01-bank-statements-to-journal-entries/",
    )
    ap.add_argument(
        "--detail",
        action="store_true",
        help="print the per-row classification under the shipped gate",
    )
    args = ap.parse_args()

    src = args.source / REF
    if not src.exists():
        print(f"source workbook not found: {src}", file=sys.stderr)
        return 1
    rows = _rows(src)
    if not graph.enabled():
        print(
            "graph disabled — set NEO4J_URI/NEO4J_USER/NEO4J_PASSWORD in .env",
            file=sys.stderr,
        )
        return 1

    try:
        from neo4j import RoutingControl

        vendors = [
            dict(rec)
            for rec in graph._DRIVER.execute_query(
                "MATCH (v:Vendor) RETURN v.name AS name, v.clean_name AS clean, "
                "v.variants AS variants",
                database_=graph._DB,
                routing_=RoutingControl.READ,
            ).records
        ]
        cands = _candidates(sorted({n for _, n, _ in rows}))
        print(f"real staging rows: {len(rows)}   seeded vendors: {len(vendors)}")
        measure_gates(rows, cands)
        measure_fold(rows, vendors)
        measure_recurrence(rows)
        if args.detail:
            print("\n=== per-row detail, shipped gate ===")
            for i, (r, pulled, golden) in enumerate(rows):
                fill = _gate(cands, pulled)
                tag = "fixture" if i < FIXTURE_ROWS else "heldout"
                print(
                    f"  {tag} r{r:<3d} {_classify(fill, golden):<32} {pulled!r}\n"
                    f"{'':14s}fill={str(fill)!r} golden={str(golden)!r}"
                )
    finally:
        graph.close_graph()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
