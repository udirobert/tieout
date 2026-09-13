#!/usr/bin/env python3
"""Seed the tieout vendor-master knowledge graph into Neo4j (Aura).

Loads the 'Vendor Master List' sheet (col A = Legal Entity Domain, col B = Vendor)
from a fixture workbook into (:Vendor) nodes keyed by a jurisdiction-stripped
vendor_key. The lineage write path then MATCHES-links resolved answer cells to
these nodes, and the GraphRAG read retrieves candidates from them. Idempotent.

Vendors are created ONLY here — harness/graph.py never MERGE-creates a Vendor,
so a written value that is not in the master links to nothing (no skeleton nodes).

Usage:
  cd research && uv run --extra graph python ../demo/seed_graph.py [init.xlsx]
"""

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "research"))
sys.path.insert(0, str(ROOT / "harness"))

import openpyxl  # noqa: E402

import graph  # noqa: E402

DEFAULT_FIXTURE = (
    ROOT
    / "demo/close-tieout/spreadsheet/close-tieout-bank-cp"
    / "1_close-tieout-bank-cp_init.xlsx"
)

_SEED_CYPHER = """
UNWIND $vendors AS v
MERGE (n:Vendor {vendor_key: v.vendor_key})
SET n.name = v.name, n.clean_name = v.clean_name,
    n.jurisdiction = v.jurisdiction, n.legal_entity_domain = v.domain
"""


def _vendor_sheet(wb):
    for ws in wb.worksheets:
        if "vendor" in ws.title.lower():
            return ws
    return None


def load_vendors(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = _vendor_sheet(wb)
        if ws is None:
            raise SystemExit(f"no 'Vendor Master List' sheet in {path}")
        vendors: list[dict] = []
        seen: set[str] = set()
        for r in range(2, ws.max_row + 1):
            domain = ws.cell(row=r, column=1).value
            name = ws.cell(row=r, column=2).value
            if not isinstance(name, str) or not name.strip():
                continue  # blank vendor rows exist in the fixture
            clean, jurisdiction = graph.split_vendor(name)
            key = graph.vendor_key(name)
            if not key or key in seen:
                continue
            seen.add(key)
            vendors.append(
                {
                    "vendor_key": key,
                    "name": name.strip(),
                    "clean_name": clean,
                    "jurisdiction": jurisdiction,
                    "domain": domain.strip() if isinstance(domain, str) else "",
                }
            )
        return vendors
    finally:
        wb.close()


def main() -> int:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_FIXTURE
    if not path.exists():
        print(f"fixture not found: {path}\nRun: python demo/build_fixtures.py")
        return 1
    if not graph.init_graph():
        print(
            "graph disabled — set NEO4J_URI/NEO4J_USER/NEO4J_PASSWORD in .env "
            "(Aura), and resume the instance if it auto-paused."
        )
        return 1
    vendors = load_vendors(path)
    ok = bool(vendors) and graph.run_write(_SEED_CYPHER, {"vendors": vendors})
    if ok:
        # Full-text is eventually consistent — wait for vendor_name_ft so a
        # seed-then-query demo never returns empty. Best-effort: proc names differ
        # across Neo4j 4.x/5.x and run_write swallows errors, so seeding never fails.
        for _await in (
            "CALL db.awaitIndex('vendor_name_ft')",
            "CALL db.index.awaitIndex('vendor_name_ft')",
            "CALL db.index.fulltext.awaitEventuallyConsistentIndexRefresh()",
        ):
            graph.run_write(_await)
    graph.close_graph()
    if not ok:
        print("seed failed (no vendors parsed, or the write errored)")
        return 1
    print(f"seeded {len(vendors)} vendors from {path.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
